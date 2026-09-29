# Copyright 2026 Nathan Gobbi
# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl).
import base64
import io
import logging

from odoo import _, api, fields, models
from odoo.exceptions import UserError

from ..tools.spreadsheet import read_rows

_logger = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:  # pragma: no cover
    _logger.debug("Cannot import openpyxl")


class MlImportMixin(models.AbstractModel):
    """Common behaviour of the spreadsheet imports.

    Concrete models define ``line_ids`` and implement ``_process_rows`` and
    ``_build_report``.
    """

    _name = "ml.import.mixin"
    _inherit = ["mail.thread"]
    _description = "Mercado Livre Spreadsheet Import"
    _order = "date_from desc, id desc"

    name = fields.Char(compute="_compute_name", store=True)
    file = fields.Binary(string="Spreadsheet", attachment=True, copy=False)
    file_name = fields.Char(copy=False)
    state = fields.Selection(
        selection=[("draft", "Draft"), ("done", "Processed")],
        default="draft",
        required=True,
        readonly=True,
        copy=False,
        tracking=True,
    )
    date_from = fields.Date(readonly=True, copy=False)
    date_to = fields.Date(readonly=True, copy=False)
    company_id = fields.Many2one(
        comodel_name="res.company",
        required=True,
        default=lambda self: self.env.company,
    )
    currency_id = fields.Many2one(related="company_id.currency_id")
    report_file = fields.Binary(readonly=True, attachment=True, copy=False)
    report_file_name = fields.Char(readonly=True, copy=False)

    def _default_name(self):
        raise NotImplementedError()

    @api.depends("date_from", "date_to")
    def _compute_name(self):
        for record in self:
            if not record.date_from:
                record.name = record._default_name()
                continue
            date_from = record.date_from.strftime("%d/%m/%Y")
            date_to = record.date_to.strftime("%d/%m/%Y")
            if date_from == date_to:
                record.name = f"{record._default_name()} {date_from}"
            else:
                record.name = f"{record._default_name()} {date_from} - {date_to}"

    def _read_file_rows(self, sheet_name=None):
        self.ensure_one()
        if not self.file:
            raise UserError(_("Upload the spreadsheet before processing it."))
        try:
            return read_rows(base64.b64decode(self.file), sheet_name=sheet_name)
        except Exception as error:
            raise UserError(
                _("Could not read the spreadsheet. Is it a .xlsx file?\n%s", error)
            ) from error

    def _process_rows(self, rows):
        """Create or update the lines of this import from the file rows."""
        raise NotImplementedError()

    def action_process(self):
        for record in self:
            if record.state != "draft":
                continue
            record._process_rows(record._read_file_rows())
            dates = record.line_ids.mapped(record._line_date_field())
            record.write(
                {
                    "state": "done",
                    "date_from": min(dates) if dates else False,
                    "date_to": max(dates) if dates else False,
                }
            )
        return True

    def action_reset_draft(self):
        for record in self:
            lines = record.line_ids
            record.write(
                {
                    "state": "draft",
                    "line_ids": [fields.Command.clear()],
                    "date_from": False,
                    "date_to": False,
                    "report_file": False,
                    "report_file_name": False,
                }
            )
            lines._unlink_orphans()
        return True

    def unlink(self):
        lines = self.mapped("line_ids")
        result = super().unlink()
        lines.exists()._unlink_orphans()
        return result

    def _line_date_field(self):
        raise NotImplementedError()

    def _generate_report(self):
        self.ensure_one()
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        file_name = self._build_report(workbook)
        buffer = io.BytesIO()
        workbook.save(buffer)
        self.write(
            {
                "report_file": base64.b64encode(buffer.getvalue()),
                "report_file_name": file_name,
            }
        )

    def _build_report(self, workbook):
        """Fill ``workbook`` with the report sheets and return the file name."""
        raise NotImplementedError()

    def action_download_report(self):
        self.ensure_one()
        # Always rebuilt, so notes edited after processing are included
        self._generate_report()
        return {
            "type": "ir.actions.act_url",
            "url": (
                f"/web/content/{self._name}/{self.id}/report_file/"
                f"{self.report_file_name}?download=true"
            ),
            "target": "self",
        }

    @staticmethod
    def _write_sheet(workbook, title, headers, rows):
        """Add a sheet with a highlighted header row and return it."""
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        font = Font(bold=True)
        for header_cell in sheet[1]:
            header_cell.fill = fill
            header_cell.font = font
        return sheet

    @staticmethod
    def _format_column(sheet, column, number_format):
        for column_cell in sheet[column][1:]:
            if isinstance(column_cell.value, int | float):
                column_cell.number_format = number_format

    @staticmethod
    def _fit_columns(sheet, max_width=40):
        for column_cells in sheet.columns:
            width = max(
                len(str(c.value)) if c.value is not None else 0 for c in column_cells
            )
            sheet.column_dimensions[column_cells[0].column_letter].width = min(
                width + 2, max_width
            )


class MlImportLineMixin(models.AbstractModel):
    _name = "ml.import.line.mixin"
    _description = "Mercado Livre Spreadsheet Import Line"

    def _unlink_orphans(self):
        """Delete the lines that no longer belong to any import."""
        self.filtered(lambda line: not line.import_ids).unlink()
