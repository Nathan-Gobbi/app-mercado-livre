# Copyright 2026 Nathan Gobbi
# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl).
"""Helpers to read the spreadsheets exported by Mercado Livre.

The functions here are plain Python (no ORM) so they can be reused and
tested in isolation.
"""

import io
import logging
import re
import unicodedata
import warnings
from datetime import date, datetime

_logger = logging.getLogger(__name__)

try:
    import openpyxl
except ImportError:  # pragma: no cover
    _logger.debug("Cannot import openpyxl")
    openpyxl = None

MONTHS_PT = {
    "janeiro": 1,
    "fevereiro": 2,
    "marco": 3,
    "abril": 4,
    "maio": 5,
    "junho": 6,
    "julho": 7,
    "agosto": 8,
    "setembro": 9,
    "outubro": 10,
    "novembro": 11,
    "dezembro": 12,
}

ML_DATE_RE = re.compile(r"(\d{1,2})\s+de\s+([a-z]+)\s+de\s+(\d{4})")


def clean_text(value):
    """Return ``value`` as a trimmed string, with non-breaking spaces removed."""
    if value is None:
        return ""
    return " ".join(str(value).replace("\xa0", " ").split())


def fold(value):
    """Normalize text for comparisons: trimmed, lowercase and without accents."""
    text = unicodedata.normalize("NFKD", clean_text(value).casefold())
    return "".join(char for char in text if not unicodedata.combining(char))


def clean_sku(value):
    """Uppercase SKU without any whitespace (``" ng 0001 "`` -> ``"NG0001"``)."""
    sku = re.sub(r"\s+", "", clean_text(value)).upper()
    return "" if sku == "NAN" else sku


def clean_code(value):
    """Convert numeric codes read as numbers (EAN, NCM) back to plain strings."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return clean_text(value)


def parse_float(value):
    """Return ``value`` as float, or ``None`` when it is empty or not a number."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return float(value)
    text = clean_text(value)
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return parse_brl_amount(text)


def parse_brl_amount(value):
    """Parse Brazilian formatted amounts such as ``"-1.234,56"``."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return float(value)
    text = clean_text(value)
    if not text:
        return None
    text = text.replace("R$", "").replace(" ", "").replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def parse_ml_date(value):
    """Parse sale dates like ``"20 de setembro de 2026 21:01 hs."``."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = ML_DATE_RE.search(fold(value))
    if not match:
        return None
    day, month_name, year = match.groups()
    month = MONTHS_PT.get(month_name)
    if not month:
        return None
    try:
        return date(int(year), month, int(day))
    except ValueError:
        return None


def parse_statement_date(value):
    """Parse account statement dates like ``"01-07-2026"``."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(clean_text(value), "%d-%m-%Y").date()
    except ValueError:
        return None


def read_rows(content, sheet_name=None, formulas=False):
    """Return the cell values of a XLSX file as a list of tuples.

    The first tuple is the first row of the sheet, so ``rows[n - 1]`` is the
    row ``n`` seen in Excel.

    :param content: raw bytes of the file
    :param sheet_name: sheet to read; the first one is used when it is missing
    :param formulas: return the formulas instead of their computed values
    """
    with warnings.catch_warnings():
        # Mercado Livre files have no default style, which openpyxl warns about
        warnings.simplefilter("ignore")
        workbook = openpyxl.load_workbook(
            io.BytesIO(content), read_only=True, data_only=not formulas
        )
    try:
        if sheet_name and sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.worksheets[0]
        return [tuple(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def find_header(rows, required):
    """Return the index of the first row that contains all ``required`` titles."""
    wanted = {fold(title) for title in required}
    for index, row in enumerate(rows):
        titles = {fold(cell) for cell in row if cell is not None}
        if wanted <= titles:
            return index
    return None


def column_index(header, candidates):
    """Find the column of the first candidate title found in ``header``.

    An exact match is preferred; otherwise a column whose title contains the
    candidate is used, like the original scripts did.
    """
    titles = [fold(cell) for cell in header]
    for candidate in candidates:
        if fold(candidate) in titles:
            return titles.index(fold(candidate))
    for candidate in candidates:
        for index, title in enumerate(titles):
            if fold(candidate) in title:
                return index
    return None


def cell(row, index):
    """Return the value at ``index`` or ``None`` when the column is missing."""
    if index is None or index >= len(row):
        return None
    return row[index]


def referenced_rows(formula, column_letter):
    """Return the rows of ``column_letter`` used by a formula.

    ``"=I139-($F$138+J139)"`` with column ``F`` returns ``[138]``.
    """
    if not isinstance(formula, str) or not formula.startswith("="):
        return []
    pattern = rf"(?<![A-Z$])\$?{column_letter}\$?(\d+)"
    return [int(row) for row in re.findall(pattern, formula)]
