# Copyright 2026 Nathan Gobbi
# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl).
"""Builders of fake spreadsheets with the same layout as the real ones."""

import base64
import io

import openpyxl

from odoo.tests.common import TransactionCase

SALES_HEADER = [
    "N.º de venda",
    "Data da venda",
    "Estado",
    "Descrição do status",
    "Unidades",
    "Receita por produtos (BRL)",
    "Total (BRL)",
    "SKU",
    "# de anúncio",
    "Título do anúncio",
    "Variação",
    "Preço unitário de venda do anúncio (BRL)",
    "Tipo de anúncio",
    "Comprador",
    "Estado",
]

STOCK_HEADER = [
    "SKU",
    "Produto",
    "Código",
    "EAN",
    "NCM",
    "Custo",
    "Tipo",
    "Venda",
    "V. Liquido",
    "Imposto",
    "Embalagem",
    "Lucro Liquido",
    "Lucro %",
    "Estoque",
    "Diferença Estoque",
]


def xlsx(rows, sheet_name="Sheet1"):
    """Return the base64 content of a XLSX file with ``rows``."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return base64.b64encode(buffer.getvalue())


def sale_row(number, date, status, units, revenue, sku, title, variation, kind):
    """A row of the sales report, with only the columns the module reads."""
    return [
        number,
        date,
        status,
        "",
        units,
        revenue,
        revenue,
        sku,
        "MLB1",
        title,
        variation,
        revenue,
        kind,
        "Fake Buyer",
        "São Paulo",
    ]


def sales_file(rows):
    """Sales report: five lines of preamble before the header, like ML."""
    preamble = [
        [None],
        ["Neste relatório, você encontra as informações das suas vendas."],
        [None],
        ["Vendas  Status das suas vendas em 23 de setembro de 2026."],
        ["Vendas", None, None, None, None, None, None, "Anúncios"],
    ]
    return xlsx(preamble + [SALES_HEADER] + rows, "Vendas BR")


def statement_file(rows, summary=("100,00", "50,00", "-30,00", "120,00")):
    return xlsx(
        [
            ["INITIAL_BALANCE", "CREDITS", "DEBITS", "FINAL_BALANCE"],
            list(summary),
            [""],
            [
                "RELEASE_DATE",
                "TRANSACTION_TYPE",
                "REFERENCE_ID",
                "TRANSACTION_NET_AMOUNT",
                "PARTIAL_BALANCE",
            ],
        ]
        + rows,
        "sheet0",
    )


def stock_file(rows):
    return xlsx([STOCK_HEADER] + rows, "Estoque")


def read_report(content):
    """Return ``{sheet title: rows}`` of a generated report."""
    workbook = openpyxl.load_workbook(io.BytesIO(base64.b64decode(content)))
    return {
        sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)]
        for sheet in workbook.worksheets
    }


class MlSalesProfitCase(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.env = cls.env(context=dict(cls.env.context, tracking_disable=True))
        cls.company = cls.env.company
        cls.company.ml_tax_rate = 6.0
        cls.product = cls.env["product.template"].create(
            {"name": "Fake Shock Kit", "default_code": "TST001", "standard_price": 100}
        )
        cls.premium = cls.env["ml.listing"].create(
            {
                "product_tmpl_id": cls.product.id,
                "listing_type": "premium",
                "price": 200.0,
                "net_amount": 150.0,
                "packaging_cost": 0.1,
            }
        )
        cls.classic = cls.env["ml.listing"].create(
            {
                "product_tmpl_id": cls.product.id,
                "listing_type": "classic",
                "price": 180.0,
                "net_amount": 140.0,
                "packaging_cost": 0.1,
            }
        )
